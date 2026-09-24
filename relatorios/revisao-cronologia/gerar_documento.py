from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

OUT=Path(__file__).resolve().parent
d=Document()
s=d.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7)
s.top_margin=Cm(1.85);s.bottom_margin=Cm(1.8);s.left_margin=Cm(2);s.right_margin=Cm(2)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
 st=d.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
 st.paragraph_format.space_after=Pt(7)
d.styles['Normal'].font.size=Pt(11)
d.styles['Normal'].paragraph_format.line_spacing=1.08
d.styles['Title'].font.size=Pt(25)
d.styles['Heading 1'].font.size=Pt(18)
d.styles['Heading 1'].paragraph_format.space_before=Pt(6)
d.styles['Heading 2'].font.size=Pt(12)
d.styles['Heading 2'].paragraph_format.space_before=Pt(10)
d.styles['Subtitle'].font.size=Pt(11)
for element in list(d.styles.element.iter(qn('w:pBdr'))):
 element.getparent().remove(element)
d.core_properties.title='Cronologia e estudo de funcionamento do projeto Imobiliária'
d.core_properties.subject='Histórico do desenvolvimento e observações de 22 de setembro de 2026'
d.core_properties.author='';d.core_properties.last_modified_by=''
foot=s.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=foot.add_run();r.font.size=Pt(9)
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');r._r.addnext(f)

def p(t): d.add_paragraph(t)
def h(t): d.add_heading(t,2)
def page(t):
 d.add_page_break();d.add_heading(t,1)
def table(headers,rows,widths):
 t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for c,w in zip(t.columns,widths):c.width=Cm(w)
 for i,txt in enumerate(headers): t.rows[0].cells[i].text=txt
 for row in rows:
  for c,txt in zip(t.add_row().cells,row):c.text=str(txt)
 for ri,row in enumerate(t.rows):
  pr=row._tr.get_or_add_trPr(); no=OxmlElement('w:cantSplit');pr.append(no)
  if ri==0:pr.append(OxmlElement('w:tblHeader'))
  for ci,c in enumerate(row.cells):
   c.width=Cm(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcp=c._tc.get_or_add_tcPr()
   sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E6EBF0' if ri==0 else 'FFFFFF');tcp.append(sh)
   borders=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
   tcp.append(borders)
   mar=OxmlElement('w:tcMar')
   for side in ['top','left','bottom','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');mar.append(e)
   tcp.append(mar)
   for pp in c.paragraphs:
    if headers[0] in ('Miniatura','Aberturas mensais') and (ci>0 or headers[0]=='Aberturas mensais'):
     pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    pp.paragraph_format.space_after=Pt(2);pp.paragraph_format.space_before=Pt(2);pp.paragraph_format.line_spacing=1
    for rr in pp.runs:rr.font.size=Pt(10);rr.bold=ri==0
 d.add_paragraph().paragraph_format.space_after=Pt(0)

d.add_paragraph('Cronologia e estudo de funcionamento do projeto Imobiliária','Title')
d.add_paragraph('Histórico de agosto e setembro de 2026\nObservações atualizadas em 22 de setembro de 2026','Subtitle')
p('O projeto reúne visualização urbana em 3D, informações de imóveis, maquetes de empreendimentos, plantas e visitas em primeira pessoa no navegador. Este documento organiza a evolução do desenvolvimento e descreve o funcionamento observado na data da revisão, incluindo a transferência de dados das miniaturas sem abrir o mapa.')
p('Na publicação consultada, o mapa de São Carlos e as três miniaturas são acessados por páginas distintas. O mapa utiliza recursos externos para detalhes urbanos; as miniaturas examinadas incorporam no HTML os dados e o código necessários aos seus modos de visualização. Essa diferença precisa acompanhar qualquer comparação de tamanho, custo ou funcionamento offline.')
h('Situação observada na revisão')
table(['Elemento','Registro em 22 de setembro'],[
('Base do projeto','README identifica a base 1.5 como desenvolvimento. Fontes organizadas em v1.5 e componentes compartilhados na raiz.'),
('Mapa publicado','São Carlos, com 11 entradas visíveis na lista de inspeção: sete casas e quatro empreendimentos.'),
('Miniaturas publicadas','Monte dos Cedros, Monte das Colinas e Wish Castanheiras, em site específico.'),
('Interação conferida','No Cedros: alternância de planta 3D, abertura do editor de móveis, visita 3D e planta 2D.'),
('Transferência medida','651.191 a 679.222 bytes gzip por HTML de miniatura, em abertura direta.'),
('Testes Node executados','113 testes: 103 aprovados e 10 reprovados. Resultado local, sem equivalência automática com a experiência publicada.')],[4,13])
h('Como interpretar as evidências')
p('As datas da cronologia descrevem registros históricos, e os números associados permanecem vinculados à versão e ao experimento de origem. As observações atuais identificam o que foi consultado ou executado nesta revisão. Os cenários de tráfego são cálculos com premissas explícitas. Nenhuma dessas categorias, isoladamente, determina qualidade visual, capacidade de escala ou prontidão comercial.')

page('Cronologia de agosto de 2026')
p('O histórico fornecido começa em 21 de agosto. A fase inicial concentrou a preparação dos dados urbanos, o desenho das edificações e a organização da carga gráfica por quarteirão. As métricas abaixo são registros do período, sem nova execução dos respectivos benchmarks. [1]')
table(['Data','Desenvolvimento registrado'],[
('21 de agosto','Inclusão de pontos de interesse no HTML. O registro inicial menciona 82 pontos em um recorte de 10 km.'),
('22 de agosto','Extração de faces do grafo viário e streaming por quarteirão no v4. O ensaio citado registra 1.281 para 238 chamadas de desenho e 14,97 para 1,88 ms por quadro.'),
('23 de agosto','No v5, separação entre página e base carregada por fetch. O HTML passou de cerca de 6,6 MB para 107 KB; o dado transferido separadamente continuava necessário.'),
('25 e 26 de agosto','Ajustes de espaço de cor, tipologias e reconstrução de lotes. No v6, dados voltaram a ser incorporados para permitir a abertura local.'),
('27 e 28 de agosto','Preparação do acervo OpenPlots, assentamento das casas nos lotes e correção da orientação de contornos. O v8 registrou redução de 10,9 para 4,4 MB em seu artefato.'),
('29 a 31 de agosto','Configuração por cidade, separação das fontes do renderizador, execução encadeada do pipeline e uso do grafo viário para complementar quadras. Foram trabalhadas outras quatro cidades.'),
('Fim de agosto','Introdução de arborização, níveis de qualidade gráfica, atualização seletiva de quadros e materiais procedurais.')],[3,14])
h('Mudanças de organização')
p('A evolução entre arquivos autocontidos e dados externos corresponde a formatos de entrega diferentes. Reduzir o HTML por separação dos dados não equivale a reduzir, na mesma proporção, o total baixado pelo visitante. O histórico dos v5 e v6 deve ser lido com essa distinção.')
p('O padrão por cidade passou a reunir projeção, caminhos e parâmetros usados no processamento. Araraquara, Ribeirão Preto, Sorocaba e São José do Rio Preto aparecem no acervo histórico. A presença desses dados não estabelece que suas páginas estejam atualmente publicadas ou tenham recebido as mesmas alterações de São Carlos.')
h('Limite das medições gráficas históricas')
p('Os ensaios relacionaram parte do tempo de quadro ao número de chamadas de desenho. Esse resultado descreve aquelas cenas e equipamentos. Não implica custo nulo para geometria adicional, nem exclui limites de GPU, memória, materiais, sombras ou processamento em outras condições.')

page('Cronologia do início de setembro de 2026')
p('Entre 1 e 8 de setembro, os registros se concentram em cadastro de interiores, extração de plantas, iluminação e bibliotecas de móveis. As comparações visuais e numéricas referem-se a ensaios específicos do período. [1]')
table(['Data','Desenvolvimento registrado'],[
('1 de setembro','No v13, entrada pela ficha do imóvel e consulta de arredores. Organização de sondas de comportamento e extração assistida de paredes a partir de plantas.'),
('3 de setembro','Uso de endereços CNEFE e consulta de lotes do Geoportal de Ribeirão Preto. Ajustes de materiais e da relação entre quadras, ruas e casas.'),
('4 de setembro','Cadastro de quatro lançamentos, representação por blocos e ajustes de ocupação dos lotes. Uma montagem de São Carlos registrou 96.168 volumes.'),
('4 e 5 de setembro','Experimentos de iluminação no v15: oclusão calculada em JavaScript, projeção de sombras e sonda de ambiente.'),
('5 de setembro','Comparação controlada entre Three.js e Unreal, com registro de diferenças de faixa dinâmica. Continuidade do desenvolvimento da visualização web em Three.js.'),
('7 e 8 de setembro','Ajustes em forros, luminárias e céu. Modelagem de móveis em Blender, variações paramétricas e mudança da posição do sol em cenas urbanas.')],[3,14])
h('Do cadastro à representação interna')
p('O histórico registra a decisão de produzir interiores a partir de plantas e informações fornecidas para cada unidade. Cômodos, alturas, aberturas e móveis alimentam a geração da geometria. A identificação dos ambientes e a conferência de escala complementam a extração por visão computacional.')
p('Um dos casos documentados apresentou divergência entre a área da planta e a ficha comercial. Essa ocorrência motivou conferências de correspondência entre anúncio, unidade e desenho. Ela não permite presumir que outras plantas apresentem a mesma divergência.')
h('Ensaios de iluminação')
p('Foram investigados sombras desativadas, superfícies coincidentes e distribuição da luz ambiente. O histórico registra faixa dinâmica de 64,1 em uma imagem Three.js e 156,8 na comparação Unreal. Esses valores dependem da câmera, iluminação, exposição e método de leitura usados no ensaio; não constituem classificação geral entre motores.')
p('Texturas procedurais e recursos calculados no navegador coexistem com dados de materiais e modelos preparados previamente. Expressões históricas como “nenhuma imagem” ou “geometria de graça” não descrevem de forma abrangente o conjunto atual de recursos.')

page('Cronologia da organização e das miniaturas')
table(['Período','Desenvolvimento registrado'],[
('10 a 13 de setembro','Registros de integração de modelos urbanos, exteriores e quintais. Parte dos detalhes passou a utilizar arquivos externos por região.'),
('14 a 18 de setembro','Modularização do renderizador e do pipeline, com testes de equivalência e comparações em navegador. O histórico descreve cerca de 80 módulos extraídos do código concentrado em app.js.'),
('18 e 19 de setembro','Reorganização do acervo e de caminhos de código. São Carlos aparece como foco do piloto. Registros de limpeza do histórico Git e retirada de caminhos antigos.'),
('19 e 20 de setembro','Desenvolvimento das transições entre mapa, interior e planta, das variantes v17 e v18 e da página independente de maquete.'),
('21 de setembro','Consolidação de fontes em v1.5. Os commits da data registram os dados reproduzíveis de São Carlos e evidências de montagem e inspeção.'),
('Estado observado em 22 de setembro','Três miniaturas publicadas com modelos de empreendimentos, plantas e visita. O Cedros expõe editor de móveis e link para o mapa. Revalidação HTTP e testes Node foram executados nesta revisão.')],[3.3,13.7])
h('Datas de implementação e datas de observação')
p('A presença de um recurso na inspeção de 22 de setembro confirma sua disponibilidade naquela consulta; não determina, por si só, o momento da implantação. Há arquivos locais fora dos commits mais recentes. Por isso, o estado publicado foi consultado separadamente do histórico Git e da documentação. [2–5]')
h('Versões e caminhos atuais')
p('“1.5” identifica a base descrita no README. “v16-moveis” é a chave da variante modular usada pelo comando de montagem. A configuração BuildConfig consultada aceita v15 e v16-moveis; diretórios v17 e v18 também existem como parte do desenvolvimento das etapas de navegação. São identificadores com funções diferentes, e sua coexistência não prova, isoladamente, divergência funcional.')
p('O mapa publicado foi acessado pela entrada do site, que encaminhou para sao-carlos-exteriores-f16b0250b705. O conteúdo HTTP recebido correspondeu ao arquivo preparado localmente com esse nome. Os HTMLs das três miniaturas também corresponderam às respectivas cópias locais examinadas.')
h('Testes e revisão registrados')
p('A execução atual dos testes Node teve 103 aprovações e 10 falhas em 113 testes, em aproximadamente 40,3 segundos. A contagem “111” do documento original foi substituída por esse resultado medido. Entre as falhas há referências a dependências ausentes no contexto dos testes; esta revisão não investigou cada causa nem corrigiu código.')
p('Os 43 testes Python aprovados pertencem ao registro de 21 de setembro e não foram reexecutados aqui. A inspeção atual do navegador complementa esses registros, mas não corresponde a um teste integral de todos os fluxos.')

page('Fontes de dados e preparação das páginas')
table(['Fonte ou entrada','Papel no projeto','Limite de interpretação'],[
('Overture Maps','Contornos de edificações e apoio à identificação de ocupação.','Um contorno pode representar apenas parte de uma construção; não equivale à divisa do lote.'),
('OpenStreetMap','Vias, nomes, pontos de interesse e atributos de edificações.','Cobertura e detalhamento variam conforme a região e a data dos dados.'),
('Cadastro municipal','Quadras, endereçamento e outras camadas usadas na composição urbana.','Disponibilidade por camada e área; extração e referência espacial precisam ser verificadas.'),
('OpenPlots','Pranchas utilizadas na reconstrução de medidas e lotes.','Uma prancha digitalizada pode exigir vetorização, escala e registro espacial.'),
('CNEFE do IBGE','Pontos de endereço utilizados no cruzamento espacial.','Endereço é indício de localização ou ocupação; não confirma sozinho a construção atual.'),
('Cadastro do imóvel','Plantas, áreas, cômodos, aberturas e elementos do empreendimento.','Dados e aproximações devem ser associados à unidade e à fonte correspondente.')],[3.4,6.3,7.3])
h('Fluxo de processamento')
p('O pipeline consulta ou lê as fontes, cruza atributos, organiza edificações por quadra, complementa limites urbanos, calcula o espaço entre quadras e vias, combina lotes de plantas e lotes sintéticos, estima ocupação, prepara volumes e detalhes e monta os arquivos de visualização. A validação é executada conforme as etapas e os comandos selecionados.')
p('O runner atual declara etapas de aquisição de fontes e etapas de transformação. Assim, a expressão “14 etapas” do texto anterior foi substituída pela descrição do fluxo: a contagem depende de incluir downloads, subetapas e rotinas opcionais. O tempo histórico de aproximadamente 485 segundos representa uma execução delimitada, sem todos os custos de aquisição pela rede. [1, 3]')
h('Coerência entre artefatos')
p('Ocupação, lotes, quadras e geometria final possuem dependências entre si. Alterar uma lista que é referenciada por índices pode exigir regenerar etapas seguintes. O projeto registra mecanismos de verificação por conteúdo para reutilizar saídas. Tempo de reconstrução e validade do cache dependem das entradas e do ambiente utilizado.')
p('Os 96.168 volumes citados na cronologia são uma contagem histórica de representação urbana. Não devem ser apresentados como quantidade atual de imóveis anunciados, unidades com interior ou construções verificadas individualmente. Na inspeção atual, a lista pública do mapa e a galeria de miniaturas têm contagens próprias.')

page('Funcionamento atual do mapa e das miniaturas')
h('Mapa urbano')
p('O mapa consultado mostra a cena de São Carlos, busca, controles de ruas, relevo e qualidade gráfica, além da lista de imóveis. A página exibe a indicação “Entorno ilustrativo”; os quatro empreendimentos na lista consultada trazem a indicação “terreno não confirmado”. A revisão observou a cena carregada, sem testar todas as ações dessa interface. [4]')
p('O módulo exterior-details.js lê tiles binários de quintais por fetch, descomprime seu conteúdo e seleciona dados conforme posição e distância. Isso é carregamento de recursos pela rede durante o uso. A montagem e exibição por proximidade na GPU também não devem ser confundidas com baixar toda a cidade novamente a cada movimento. [3]')
h('Miniaturas independentes')
p('Os três HTMLs publicados incorporam biblioteca de renderização, dados do empreendimento, geometria e componentes de planta. O gerador inclui exportações preparadas em Blender e código Three.js para exibição e interação. Projetos Blender e modelos GLB são oferecidos separadamente na galeria; não precisam ser baixados para abrir o HTML da miniatura. [3, 5]')
table(['Função','Funcionamento identificado'],[
('Prédio e conjunto','Geometria carregada no documento e controles executados no navegador.'),
('Planta 2D','Desenho da planta a partir dos dados incorporados.'),
('Planta 3D e visita','Construção e renderização da cena interna no aparelho do visitante.'),
('Móveis no Cedros','Editor disponível na planta 3D; o código salva alterações em localStorage, no navegador.'),
('Ficha','Informações incorporadas na geração da página; não é demonstrada consulta ao banco por abertura.'),
('Ver mapa','Navegação para outro documento; fora da contabilização de tráfego das miniaturas.')],[4,13])
h('Divisão do trabalho')
p('Blender e o pipeline preparam artefatos antes da publicação. O Hosting entrega os arquivos solicitados. O navegador processa a geometria e executa os modos de visualização. A análise do código das miniaturas não identificou necessidade de renderização no servidor ou de consulta ao Firestore para essas funções.')
p('Isso descreve a arquitetura examinada, não uma medição de CPU ou memória do provedor. O arquivo transferido, a memória ocupada pela cena e o tempo por quadro são grandezas distintas. A existência de regras e dependências Firebase no repositório também não comprova que cada página publicada as utilize.')

page('Medições atuais de transferência e cache')
p('Em 22 de setembro foram enviados GETs às páginas publicadas com Accept-Encoding: gzip. O corpo recebido foi contado antes da descompressão e comparado, após descompressão, com os arquivos locais. As três miniaturas apresentaram correspondência de conteúdo. KB, MB e GB neste estudo usam unidades decimais. [5]')
table(['Miniatura','HTML em bytes','Corpo gzip em bytes','Redução'],[
('Monte dos Cedros','3.361.828','673.852','80,0%'),
('Monte das Colinas','3.316.385','651.191','80,4%'),
('Wish Castanheiras','4.001.617','679.222','83,0%')],[5,4.2,4.8,3])
p('Os valores medem o corpo HTTP do documento. Excluem cabeçalhos, estabelecimento de conexão e demais custos do protocolo. Outras negociações de compressão podem gerar tamanhos diferentes. A medição não é de tempo até interação, memória RAM ou desempenho gráfico.')
h('Primeira abertura e retorno')
p('As miniaturas responderam com Cache-Control: public, max-age=0, must-revalidate e ETag. Em uma segunda requisição com If-None-Match, cada página respondeu HTTP 304, sem corpo. Isso confirma a revalidação nessa condição de teste. O reaproveitamento exige uma cópia local válida; atualizações, cache ausente ou políticas do navegador podem causar nova transferência integral.')
p('A navegação observada no Cedros permitiu acessar planta 3D, editor, visita e planta 2D. A ausência de chamadas adicionais de aplicação para essas funções decorre da inspeção do código incorporado. Não foi gravada uma captura completa de rede de todos os comandos e de todas as unidades.')
h('Galeria e armazenamento')
p('A entrada da galeria transferiu 1.600 bytes gzip e referencia três PNGs, cujos arquivos locais publicados somam 2.583.491 bytes. Assim, galeria seguida de uma miniatura representa aproximadamente 3,24 a 3,26 MB de corpo, se os PNGs forem entregues nesses tamanhos e não houver cache. As imagens não foram baixadas novamente nesta medição.')
p('Os 13 arquivos preparados para o site de miniaturas somam 18.821.140 bytes locais. Esse total inclui imagens, GLBs e projetos Blender; apenas os três HTMLs somam 10.679.830 bytes. O armazenamento faturado pode também incluir versões anteriores e não foi consultado.')
h('Mapa como referência separada')
p('O documento de entrada do mapa versionado transferiu 9.801.638 bytes gzip e continha 13.957.621 bytes após a descompressão HTTP. Esse tamanho é do artefato entregue, não da cena expandida em memória nem da sessão inteira com tiles. O cabeçalho recebido foi public, max-age=31536000, immutable. Os números históricos de 32,9 MB e 14,0 MB não são uma medição atual de sessão.')

page('Cenários de tráfego e custo')
p('Para simulação, considera-se 0,70 MB por abertura direta de miniatura com transferência integral. O valor arredonda as medições gzip desta revisão. “Abertura” significa carregamento do documento; alternar os modos dentro dele não representa uma nova abertura.')
table(['Aberturas mensais','Corpo estimado','Transferência estimada no Blaze'],[
('1.000','0,7 GB','US$ 0,00'),('10.000','7 GB','US$ 0,00'),('100.000','70 GB','US$ 9,00'),('1.000.000','700 GB','US$ 103,50')],[5,4,8])
p('A referência de preço consultada para Firebase Hosting informa franquia de 10 GB/mês e US$ 0,15 por GB excedente no Blaze. A tabela supõe essa franquia inteira disponível para as miniaturas. As cotas são compartilhadas no projeto, e entregas a partir da CDN também entram no tráfego. No Spark, o excesso tem tratamento de limite de serviço, não de cobrança automática. [6]')
p('Cálculo adotado: tráfego em GB = aberturas × 0,70 ÷ 1.000. Custo em dólares = máximo entre zero e (tráfego − 10) × 0,15. São estimativas de transferência, sem tributos, câmbio, outros serviços ou armazenamento excedente. O plano contratado e a utilização real do projeto não foram consultados.')
h('Como o percurso altera o resultado')
p('O percurso pela galeria acrescenta o carregamento das imagens. Usando 3,30 MB como arredondamento para galeria mais uma miniatura, 100 mil percursos completos sem cache resultariam em aproximadamente 330 GB e US$ 48 de transferência, sob as mesmas premissas. Downloads de GLB ou Blender precisariam ser somados à parte.')
p('Uma revisita com revalidação 304 transfere cabeçalhos, mas pode dispensar o corpo do HTML. Uma nova versão do documento volta a exigir seu conteúdo. Portanto, o tráfego mensal depende da proporção de primeiras visitas, retornos, atualizações e percursos realizados, além da quantidade de pessoas.')
h('Simultaneidade e banda')
p('Cem novas aberturas por segundo, a 0,70 MB cada, corresponderiam a aproximadamente 70 MB/s ou 560 Mbit/s de corpo transferido. Trata-se de uma conversão aritmética, não de capacidade comprovada do Hosting ou de um servidor próprio. Cem pessoas com a página já carregada não têm esse mesmo perfil de transferência.')
h('O que estes cálculos permitem comparar')
p('Os cenários permitem comparar entrada direta, passagem pela galeria e aproveitamento de cache mantendo as premissas visíveis. Não determinam a capacidade de usuários simultâneos, a experiência em conexão móvel, o consumo de bateria ou o desempenho em diferentes aparelhos. Esses resultados requerem medições próprias.')

page('Alcance da revisão e referências')
h('Verificações realizadas em 22 de setembro')
p('Foram lidos o documento de origem, registros do projeto, configuração de publicação, gerador das miniaturas e módulos relacionados. Foram consultados o histórico Git recente, os HTMLs publicados e os cabeçalhos HTTP. O mapa foi observado após carregar; no Cedros foram acessados os modos e o painel de móveis descritos neste relatório. A suíte Node foi executada localmente.')
h('Aspectos que permanecem delimitados')
p('Os experimentos históricos não foram reexecutados. Não houve ensaio de carga concorrente, medição de RAM, GPU, bateria ou FPS em celular físico, consulta à fatura ou verificação das regras efetivamente implantadas no Firestore. As observações locais dessas regras não devem ser apresentadas como descrição confirmada das permissões em produção.')
p('Os registros consultados não documentam, neste recorte, uma avaliação de uso com participantes que permita concluir sobre compreensão da planta, intenção de contato ou resultado comercial. A ausência desse registro não equivale a afirmar que nenhuma avaliação tenha ocorrido fora do material examinado.')
p('O resultado 103/113 dos testes Node e a inspeção dos fluxos selecionados têm alcances diferentes. Para associar uma falha de teste a um comportamento publicado, seria necessário reproduzir o caso correspondente. Esta revisão registra evidências e condições de observação, sem atribuir aprovação ou reprovação global ao projeto.')
h('Referências')
p('[1] Cronologia do projeto Imobiliária.docx, documento fornecido para revisão, datado de 22/09/2026. Fonte dos eventos e métricas históricos, salvo indicação de verificação atual.')
p('[2] README.md, REVISAO-1.5.md, v1.5/LEIA-ME.md e histórico Git local. Commits consultados incluem 46c1c0c, de 20/09, e da01e5d e a3502ce, de 21/09/2026.')
p('[3] Código local consultado em 22/09: pipeline/rodar.py; pipeline/build/config.py; v1.5/miniaturas/pagina_maquete.py, editor_cedros.js e preparar_publicacao.py; v1.5/renderizador-v16-moveis/exterior-details.js; firebase.json; firebase.miniaturas.json; firebase/firestore.rules.')
p('[4] Mapa publicado: https://imobilaria-deccb.web.app/ — entrada observada para /mapa/sao-carlos-exteriores-f16b0250b705. Consulta em 22/09/2026.')
p('[5] Miniaturas publicadas: https://imobilaria-deccb-miniaturas.web.app/ — maquete-monte-dos-cedros-37.html, maquete-monte-das-colinas-39.html e maquete-wish-castanheiras-58.html. Medições HTTP em 22/09/2026.')
p('[6] Firebase Hosting — Learn about usage levels, quotas, and pricing for Hosting. https://firebase.google.com/docs/hosting/usage-quotas-pricing. Consulta em 22/09/2026.')

out=OUT/'Cronologia e estudo do projeto Imobiliária revisados.docx'
d.save(out)
print(out)
