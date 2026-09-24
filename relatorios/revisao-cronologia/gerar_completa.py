from pathlib import Path
BASE=Path(__file__).resolve().parent
# Reutiliza somente as definições de formatação, sem executar a autoria anterior.
setup=(BASE/'gerar_documento.py').read_text(encoding='utf8').split("d.add_paragraph('Cronologia e estudo de funcionamento do projeto Imobiliária','Title')")[0]
exec(compile(setup,str(BASE/'gerar_documento.py'),'exec'))
d.core_properties.title='Cronologia e estudo detalhado do projeto Imobiliária'
d.core_properties.subject='Versão extensa com histórico e observações de 22 de setembro de 2026'
for name in ['TOC 1','TOC Heading']:
 if name not in d.styles:
  from docx.enum.style import WD_STYLE_TYPE
  d.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
 d.styles[name].font.name='Calibri';d.styles[name].font.size=Pt(10)
 d.styles[name].font.color.rgb=RGBColor(0,0,0)
 d.styles[name].paragraph_format.space_after=Pt(1)
 d.styles[name].paragraph_format.line_spacing=1

lines=(BASE/'versao-completa.md').read_text(encoding='utf8').splitlines()
intro=True;i=0
while i<len(lines):
 line=lines[i].strip()
 if not line:i+=1;continue
 if line.startswith('# '):d.add_paragraph(line[2:],'Title')
 elif line.startswith('## '):
  if intro:
   d.add_paragraph('Sumário','Title').paragraph_format.page_break_before=True
   q=d.add_paragraph();r=q.add_run();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'TOC \\o "1-1" \\h \\z \\u');r._r.addnext(fld)
   intro=False
  d.add_heading(line[3:],1).paragraph_format.page_break_before=True
 elif line.startswith('### '):d.add_heading(line[4:],2)
 elif line.startswith('|'):
  rows=[]
  while i<len(lines) and lines[i].strip().startswith('|'):
   row=[v.strip() for v in lines[i].strip().strip('|').split('|')]
   if not all(set(v)<=set('-: ') for v in row):rows.append(row)
   i+=1
  widths={2:[12,5],3:[5,4,8],4:[5,4.2,4.5,3.3]}[len(rows[0])]
  table(rows[0],rows[1:],widths)
  for row in d.tables[-1].rows:
   for ci,cell in enumerate(row.cells):
    if ci>0 or rows[0][0]=='Aberturas por mês':
     for para in cell.paragraphs:para.alignment=WD_ALIGN_PARAGRAPH.CENTER
  continue
 elif line.startswith('Histórico do desenvolvimento'):d.add_paragraph(line,'Subtitle')
 else:d.add_paragraph(line)
 i+=1

out=BASE/'Cronologia e estudo detalhado do projeto Imobiliária.docx'
d.save(out)
print(out)
